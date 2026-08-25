#!/usr/bin/env python3
"""Generate multi-symbol-regression test sets from a physicsMDSR workbook.

Each workbook row becomes an ``<output-root>/<ID>/`` directory.  For a WASS
label such as ``WASS_0``, the script pairs the eight vectors in
``paraValue_WASS_0`` with the eight range groups in ``varRange_WASS_0`` and
writes ``0.csv`` through ``7.csv``.

Example:
    python generate_mdsr_testsets.py physicsMDSR_WASS_varRanges.xlsx \
        --wass WASS_0 --samples 5000 --output-root WASS_0_testsets
"""

from __future__ import annotations

import argparse
import ast
import keyword
import math
import os
import re
import sys
import zlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

import numpy as np
from openpyxl import load_workbook


# Change this value (or use --wass) to reuse the program for WASS_025/WASS_04.
WASS_LABEL = "WASS_0"
DEFAULT_SAMPLES = 5000
DEFAULT_SEED = 20260824
DEFAULT_SHEET = "Sampling design"
EXPECTED_VECTOR_COUNT = 8


FUNCTIONS: dict[str, Any] = {
    "sqrt": np.sqrt,
    "exp": np.exp,
    "log": np.log,
    "ln": np.log,
    "log10": np.log10,
    "sin": np.sin,
    "cos": np.cos,
    "tan": np.tan,
    "asin": np.arcsin,
    "acos": np.arccos,
    "atan": np.arctan,
    "arcsin": np.arcsin,
    "arccos": np.arccos,
    "arctan": np.arctan,
    "sinh": np.sinh,
    "cosh": np.cosh,
    "tanh": np.tanh,
    "asinh": np.arcsinh,
    "acosh": np.arccosh,
    "atanh": np.arctanh,
    "arcsinh": np.arcsinh,
    "arccosh": np.arccosh,
    "arctanh": np.arctanh,
    "abs": np.abs,
}

BUILTIN_CONSTANTS = {"pi": math.pi, "e": math.e}
NAME_RE = re.compile(r"^[A-Za-z_]\w*$")
PARAMETER_RE = re.compile(r"([A-Za-z_]\w*)\s*=\s*\[")


class WorkbookDataError(ValueError):
    """Raised when a workbook cell does not satisfy the expected schema."""


@dataclass(frozen=True)
class RowSpec:
    row_number: int
    model_id: str
    formula: str
    target: str
    variables: tuple[str, ...]
    parameter_names: tuple[str, ...]
    constants: Mapping[str, float]
    parameter_vectors: tuple[tuple[float, ...], ...]
    variable_range_groups: tuple[tuple[tuple[float, float], ...], ...]
    compiled_formula: Any
    symbol_aliases: Mapping[str, str]


class FormulaValidator(ast.NodeVisitor):
    """Allow arithmetic expressions and a small set of NumPy functions only."""

    _binary_ops = (ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow, ast.Mod)
    _unary_ops = (ast.UAdd, ast.USub)

    def __init__(self, allowed_symbols: set[str]) -> None:
        self.allowed_symbols = allowed_symbols

    def visit_Expression(self, node: ast.Expression) -> None:
        self.visit(node.body)

    def visit_BinOp(self, node: ast.BinOp) -> None:
        if not isinstance(node.op, self._binary_ops):
            raise WorkbookDataError(f"不允许的二元运算符: {type(node.op).__name__}")
        self.visit(node.left)
        self.visit(node.right)

    def visit_UnaryOp(self, node: ast.UnaryOp) -> None:
        if not isinstance(node.op, self._unary_ops):
            raise WorkbookDataError(f"不允许的一元运算符: {type(node.op).__name__}")
        self.visit(node.operand)

    def visit_Call(self, node: ast.Call) -> None:
        if not isinstance(node.func, ast.Name) or node.func.id not in FUNCTIONS:
            raise WorkbookDataError("公式只允许调用受支持的数学函数")
        if node.keywords:
            raise WorkbookDataError("公式中的函数调用不允许关键字参数")
        for arg in node.args:
            self.visit(arg)

    def visit_Name(self, node: ast.Name) -> None:
        if node.id not in self.allowed_symbols and node.id not in FUNCTIONS:
            raise WorkbookDataError(f"公式含未定义符号: {node.id}")

    def visit_Constant(self, node: ast.Constant) -> None:
        if isinstance(node.value, bool) or not isinstance(node.value, (int, float)):
            raise WorkbookDataError("公式只允许数值常量")

    def generic_visit(self, node: ast.AST) -> None:
        raise WorkbookDataError(f"公式含不允许的语法: {type(node).__name__}")


def normalize_wass_label(value: str) -> str:
    value = value.strip()
    if not value:
        raise argparse.ArgumentTypeError("WASS 标签不能为空")
    if value.upper().startswith("WASS_"):
        return "WASS_" + value[5:]
    return "WASS_" + value


def split_names(cell: Any, *, field: str) -> tuple[str, ...]:
    names = tuple(part.strip() for part in re.split(r"[,;]", str(cell)) if part.strip())
    if not names:
        raise WorkbookDataError(f"{field} 为空")
    if len(set(names)) != len(names) or any(not NAME_RE.fullmatch(x) for x in names):
        raise WorkbookDataError(f"{field} 含重复或非法名称: {names}")
    return names


def parse_parameter_names(cell: Any) -> tuple[str, ...]:
    names = tuple(PARAMETER_RE.findall(str(cell)))
    if not names:
        raise WorkbookDataError("ParameterRange 中未找到参数名称")
    if len(set(names)) != len(names):
        raise WorkbookDataError(f"ParameterRange 含重复参数: {names}")
    return names


def parse_constants(cell: Any) -> dict[str, float]:
    if cell is None or str(cell).strip().lower() in {"", "(none)", "none", "nan"}:
        return {}
    result: dict[str, float] = {}
    for item in str(cell).split(";"):
        if "=" not in item:
            raise WorkbookDataError(f"无法解析常量: {item!r}")
        name, raw_value = (x.strip() for x in item.split("=", 1))
        if not NAME_RE.fullmatch(name) or name in result:
            raise WorkbookDataError(f"常量名称非法或重复: {name!r}")
        try:
            value = float(raw_value)
        except ValueError as exc:
            raise WorkbookDataError(f"常量 {name} 不是数值: {raw_value!r}") from exc
        if not math.isfinite(value):
            raise WorkbookDataError(f"常量 {name} 不是有限数")
        result[name] = value
    return result


def parse_list_cell(cell: Any, field: str) -> list[Any]:
    if isinstance(cell, (list, tuple)):
        value = cell
    else:
        try:
            value = ast.literal_eval(str(cell).strip())
        except (SyntaxError, ValueError) as exc:
            raise WorkbookDataError(f"{field} 不是合法的向量列表") from exc
    if not isinstance(value, (list, tuple)):
        raise WorkbookDataError(f"{field} 必须是列表")
    return list(value)


def as_finite_float(value: Any, field: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise WorkbookDataError(f"{field} 含非数值: {value!r}") from exc
    if not math.isfinite(number):
        raise WorkbookDataError(f"{field} 含非有限数: {value!r}")
    return number


def parse_parameter_vectors(cell: Any, expected_dimension: int, field: str) -> tuple[tuple[float, ...], ...]:
    outer = parse_list_cell(cell, field)
    if len(outer) != EXPECTED_VECTOR_COUNT:
        raise WorkbookDataError(f"{field} 应含 {EXPECTED_VECTOR_COUNT} 组，实际为 {len(outer)} 组")
    result = []
    for i, vector in enumerate(outer):
        if not isinstance(vector, (list, tuple)) or len(vector) != expected_dimension:
            raise WorkbookDataError(f"{field} 第 {i} 组维度应为 {expected_dimension}")
        result.append(tuple(as_finite_float(x, field) for x in vector))
    return tuple(result)


def parse_variable_ranges(
    cell: Any, expected_dimension: int, field: str
) -> tuple[tuple[tuple[float, float], ...], ...]:
    outer = parse_list_cell(cell, field)
    if len(outer) != EXPECTED_VECTOR_COUNT:
        raise WorkbookDataError(f"{field} 应含 {EXPECTED_VECTOR_COUNT} 组，实际为 {len(outer)} 组")
    result = []
    for group_index, group in enumerate(outer):
        if not isinstance(group, (list, tuple)) or len(group) != expected_dimension:
            raise WorkbookDataError(f"{field} 第 {group_index} 组变量数应为 {expected_dimension}")
        parsed_group = []
        for variable_index, interval in enumerate(group):
            if not isinstance(interval, (list, tuple)) or len(interval) != 2:
                raise WorkbookDataError(
                    f"{field} 第 {group_index} 组第 {variable_index} 个范围必须为 [下限, 上限]"
                )
            low = as_finite_float(interval[0], field)
            high = as_finite_float(interval[1], field)
            if not low < high:
                raise WorkbookDataError(f"{field} 范围必须满足下限 < 上限: {interval}")
            parsed_group.append((low, high))
        result.append(tuple(parsed_group))
    return tuple(result)


def compile_formula(
    formula_cell: Any, allowed_symbols: set[str]
) -> tuple[str, Any, Mapping[str, str]]:
    formula = str(formula_cell).strip().replace("^", "**")
    if not formula:
        raise WorkbookDataError("GenerationFormula 为空")
    # Physics symbols may be Python keywords (the workbook uses ``lambda``).
    aliases = {
        name: f"_symbol_{name}"
        for name in allowed_symbols
        if keyword.iskeyword(name)
    }
    parse_formula = formula
    for original, alias in aliases.items():
        parse_formula = re.sub(rf"\b{re.escape(original)}\b", alias, parse_formula)
    parse_symbols = {aliases.get(name, name) for name in allowed_symbols}
    try:
        tree = ast.parse(parse_formula, mode="eval")
    except SyntaxError as exc:
        raise WorkbookDataError(f"GenerationFormula 语法错误: {formula}") from exc
    FormulaValidator(parse_symbols).visit(tree)
    return formula, compile(tree, "<GenerationFormula>", "eval"), aliases


def load_specs(workbook_path: Path, sheet_name: str, wass_label: str) -> list[RowSpec]:
    parameter_column = f"paraValue_{wass_label}"
    range_column = f"varRange_{wass_label}"
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise WorkbookDataError(f"找不到工作表 {sheet_name!r}; 可用工作表: {workbook.sheetnames}")
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(x).strip() if x is not None else "" for x in next(rows)]
    required = {
        "ID", "GenerationFormula", "Target", "IndependentVars", "ParameterRange",
        "FixedConstantValues", parameter_column, range_column,
    }
    missing = sorted(required.difference(headers))
    if missing:
        raise WorkbookDataError(f"工作表缺少列: {', '.join(missing)}")
    column = {name: headers.index(name) for name in required}

    specs: list[RowSpec] = []
    seen_ids: set[str] = set()
    for row_number, row in enumerate(rows, start=2):
        if all(value is None for value in row):
            continue
        try:
            model_id = str(row[column["ID"]]).strip()
            if not model_id or model_id in {".", ".."} or Path(model_id).name != model_id:
                raise WorkbookDataError(f"ID 不能安全地用作文件夹名: {model_id!r}")
            if model_id in seen_ids:
                raise WorkbookDataError(f"ID 重复: {model_id}")
            seen_ids.add(model_id)

            variables = split_names(row[column["IndependentVars"]], field="IndependentVars")
            target = str(row[column["Target"]]).strip()
            if not NAME_RE.fullmatch(target):
                raise WorkbookDataError(f"Target 名称非法: {target!r}")
            parameter_names = parse_parameter_names(row[column["ParameterRange"]])
            constants = parse_constants(row[column["FixedConstantValues"]])
            all_names = list(variables) + list(parameter_names) + list(constants)
            if len(all_names) != len(set(all_names)):
                raise WorkbookDataError("变量、参数和常量名称存在冲突")

            parameter_vectors = parse_parameter_vectors(
                row[column[parameter_column]], len(parameter_names), parameter_column
            )
            variable_ranges = parse_variable_ranges(
                row[column[range_column]], len(variables), range_column
            )
            allowed_symbols = set(all_names) | set(BUILTIN_CONSTANTS)
            formula, compiled, aliases = compile_formula(
                row[column["GenerationFormula"]], allowed_symbols
            )
            specs.append(
                RowSpec(
                    row_number=row_number,
                    model_id=model_id,
                    formula=formula,
                    target=target,
                    variables=variables,
                    parameter_names=parameter_names,
                    constants=constants,
                    parameter_vectors=parameter_vectors,
                    variable_range_groups=variable_ranges,
                    compiled_formula=compiled,
                    symbol_aliases=aliases,
                )
            )
        except WorkbookDataError as exc:
            raise WorkbookDataError(f"Excel 第 {row_number} 行: {exc}") from exc
    if not specs:
        raise WorkbookDataError("工作表中没有可处理的数据行")
    return specs


def evaluate_formula(spec: RowSpec, variables: Mapping[str, np.ndarray], parameters: Sequence[float]) -> np.ndarray:
    environment: dict[str, Any] = {}
    environment.update(FUNCTIONS)
    environment.update(BUILTIN_CONSTANTS)
    environment.update(spec.constants)  # Workbook constants intentionally override pi/e.
    environment.update(zip(spec.parameter_names, parameters))
    environment.update(variables)
    for original, alias in spec.symbol_aliases.items():
        environment[alias] = environment[original]
    with np.errstate(all="ignore"):
        result = eval(spec.compiled_formula, {"__builtins__": {}}, environment)
    values = np.asarray(result)
    sample_count = len(next(iter(variables.values())))
    if values.ndim == 0:
        values = np.full(sample_count, values.item())
    else:
        try:
            values = np.broadcast_to(values, (sample_count,))
        except ValueError as exc:
            raise WorkbookDataError(f"{spec.model_id} 的公式输出形状不是一维样本") from exc
    if np.iscomplexobj(values):
        values = np.where(np.abs(values.imag) <= 1e-12, values.real, np.nan)
    try:
        return values.astype(float, copy=False)
    except (TypeError, ValueError) as exc:
        raise WorkbookDataError(f"{spec.model_id} 的公式输出无法转换为实数") from exc


def generate_dataset(
    spec: RowSpec,
    parameter_index: int,
    sample_count: int,
    seed: int,
    max_draw_multiplier: int,
) -> np.ndarray:
    parameters = spec.parameter_vectors[parameter_index]
    ranges = spec.variable_range_groups[parameter_index]
    stable_id = zlib.crc32(spec.model_id.encode("utf-8"))
    rng = np.random.default_rng(np.random.SeedSequence([seed, stable_id, parameter_index]))
    accepted: list[np.ndarray] = []
    accepted_count = 0
    drawn_count = 0
    max_draws = max(sample_count, sample_count * max_draw_multiplier)

    while accepted_count < sample_count and drawn_count < max_draws:
        remaining = sample_count - accepted_count
        batch_size = min(max(remaining * 2, 1024), max_draws - drawn_count)
        columns = {
            name: rng.uniform(low, high, size=batch_size)
            for name, (low, high) in zip(spec.variables, ranges)
        }
        target = evaluate_formula(spec, columns, parameters)
        matrix = np.column_stack([*(columns[name] for name in spec.variables), target])
        valid = np.all(np.isfinite(matrix), axis=1)
        if np.any(valid):
            chunk = matrix[valid][:remaining]
            accepted.append(chunk)
            accepted_count += len(chunk)
        drawn_count += batch_size

    if accepted_count < sample_count:
        raise WorkbookDataError(
            f"{spec.model_id}/{parameter_index}: 仅获得 {accepted_count}/{sample_count} 个有限实数样本；"
            "请检查变量范围和公式定义域"
        )
    return np.vstack(accepted)


def preflight_destinations(specs: Sequence[RowSpec], output_root: Path, overwrite: bool) -> None:
    if output_root.exists() and not output_root.is_dir():
        raise FileExistsError(f"输出根路径不是文件夹: {output_root}")
    if overwrite:
        return
    existing = [
        output_root / spec.model_id / f"{index}.csv"
        for spec in specs
        for index in range(EXPECTED_VECTOR_COUNT)
        if (output_root / spec.model_id / f"{index}.csv").exists()
    ]
    if existing:
        preview = ", ".join(str(path) for path in existing[:3])
        raise FileExistsError(f"已有输出文件（例如 {preview}）；如需覆盖请加 --overwrite")


def write_dataset(path: Path, data: np.ndarray, column_names: Sequence[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    try:
        np.savetxt(
            temporary,
            data,
            delimiter=",",
            header=",".join(column_names),
            comments="",
            fmt="%.12g",
        )
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="从 physicsMDSR 工作簿生成多符号回归测试集")
    parser.add_argument("workbook", type=Path, help="输入 .xlsx 文件")
    parser.add_argument("--wass", type=normalize_wass_label, default=WASS_LABEL, help="WASS 标签，默认 WASS_0")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help=f"工作表名称，默认 {DEFAULT_SHEET!r}")
    parser.add_argument("--samples", type=int, default=DEFAULT_SAMPLES, help="每个 CSV 的样本数")
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED, help="可复现的随机种子")
    parser.add_argument("--output-root", type=Path, help="输出根文件夹；默认位于输入文件旁")
    parser.add_argument("--overwrite", action="store_true", help="覆盖已存在的 0.csv～7.csv")
    parser.add_argument("--validate-only", action="store_true", help="只校验工作簿，不生成 CSV")
    parser.add_argument(
        "--max-draw-multiplier", type=int, default=50,
        help="为剔除公式定义域外样本允许的最大抽样倍数（默认 50）",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if args.samples <= 0:
        raise ValueError("--samples 必须大于 0")
    if args.max_draw_multiplier <= 0:
        raise ValueError("--max-draw-multiplier 必须大于 0")
    workbook_path = args.workbook.expanduser().resolve()
    if not workbook_path.is_file():
        raise FileNotFoundError(f"找不到输入文件: {workbook_path}")
    output_root = (
        args.output_root.expanduser().resolve()
        if args.output_root
        else workbook_path.with_name(f"{workbook_path.stem}_{args.wass}_testsets")
    )

    specs = load_specs(workbook_path, args.sheet, args.wass)
    print(
        f"校验通过：{len(specs)} 行；使用 paraValue_{args.wass} / varRange_{args.wass}；"
        f"每行 {EXPECTED_VECTOR_COUNT} 组。"
    )
    if args.validate_only:
        return 0

    preflight_destinations(specs, output_root, args.overwrite)
    total_files = len(specs) * EXPECTED_VECTOR_COUNT
    completed = 0
    for spec in specs:
        for index in range(EXPECTED_VECTOR_COUNT):
            data = generate_dataset(
                spec, index, args.samples, args.seed, args.max_draw_multiplier
            )
            write_dataset(
                output_root / spec.model_id / f"{index}.csv",
                data,
                (*spec.variables, spec.target),
            )
            completed += 1
        print(f"[{completed:>3}/{total_files}] {spec.model_id} 完成")
    print(f"生成完成：{total_files} 个 CSV，输出目录：{output_root}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (WorkbookDataError, FileNotFoundError, FileExistsError, ValueError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        raise SystemExit(2)
