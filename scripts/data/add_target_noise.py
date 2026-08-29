#!/usr/bin/env python3
"""Create CSV datasets with reproducible Gaussian noise on the target.

For every ID in the metadata workbook, the program processes a configurable
inclusive range of directly contained CSV files. The default range is
DATA_FOLDER/ID/0.csv through DATA_FOLDER/ID/7.csv. Use ``--dataset-start`` and
``--dataset-end`` to change it, for example ``--dataset-end 15`` for 0..15.csv.
Other filenames and CSVs inside nested directories are ignored. For each
requested noise amplitude n, it writes a corresponding copy below
NOISE_FOLDER/<noise-value>/ID. Output files are renumbered from zero: the input
file at ``--dataset-start`` becomes 0.csv, the next becomes 1.csv, and so on.

The default noise model is relative Gaussian noise:

    noisy_target = target * (1 + n * Z),  Z ~ Normal(0, 1)

Thus n is the one-standard-deviation relative amplitude: n=0.03 means 3%
Gaussian noise. Two additional modes are available:

    dataset-std: noisy_target = target + n * std(target_column) * Z
    absolute:    noisy_target = target + n * Z

Use ``--seed`` to reproduce exactly the same noise. Each output file gets a
stable, independent random stream, so adding or removing another input file
does not change its generated values.

The target column name is read from the workbook's ``Target`` column.  Thus an
ID whose Target value is ``v`` modifies the CSV column named ``v``.

Example output directory names:
    0.01 -> 001
    0.03 -> 003
    0.05 -> 005
    0.1  -> 01
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import math
import random
import sys
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - depends on the user's machine
    raise SystemExit(
        "Missing dependency: openpyxl. Install it with: pip install openpyxl"
    ) from exc


DEFAULT_NOISES = "n1=0.01,n2=0.03,n3=0.05,n4=0.1"
DEFAULT_DATASET_START = 0
DEFAULT_DATASET_END = 7
ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")


@dataclass(frozen=True)
class CsvInfo:
    source: Path
    relative_to_id: Path
    encoding: str
    dialect: type[csv.Dialect]
    target_column: str


def parse_noise_list(text: str) -> list[Decimal]:
    """Parse either 'n1=0.01,n2=0.03' or '0.01,0.03'."""
    values: list[Decimal] = []
    seen: set[Decimal] = set()

    for item in text.split(","):
        item = item.strip()
        if not item:
            continue
        raw_value = item.split("=", 1)[-1].strip()
        try:
            value = Decimal(raw_value)
        except InvalidOperation as exc:
            raise ValueError(f"Invalid noise value: {raw_value!r}") from exc
        if not value.is_finite() or value < 0:
            raise ValueError(f"Noise must be a finite non-negative number: {raw_value!r}")
        if value in seen:
            raise ValueError(f"Duplicate noise value: {raw_value!r}")
        seen.add(value)
        values.append(value)

    if not values:
        raise ValueError("The noise list is empty.")
    return values


def noise_directory_name(value: Decimal) -> str:
    """Convert 0.01 to '001', 0.03 to '003', and 0.1 to '01'."""
    text = format(value.normalize(), "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    if text.startswith("0."):
        return "0" + text[2:]
    return text.replace(".", "")


def load_id_targets(workbook_path: Path, sheet_name: str) -> dict[str, str]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        workbook.close()
        raise ValueError(f"Sheet {sheet_name!r} not found. Available sheets: {available}")

    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        workbook.close()
        raise ValueError("The metadata sheet is empty.") from exc

    headers = [str(value).strip() if value is not None else "" for value in header_row]
    try:
        id_index = headers.index("ID")
        target_index = headers.index("Target")
    except ValueError as exc:
        workbook.close()
        raise ValueError("The metadata sheet must contain columns named ID and Target.") from exc

    id_targets: dict[str, str] = {}
    for excel_row, row in enumerate(rows, start=2):
        raw_id = row[id_index] if id_index < len(row) else None
        raw_target = row[target_index] if target_index < len(row) else None
        if raw_id is None or str(raw_id).strip() == "":
            continue
        item_id = str(raw_id).strip()
        target = "" if raw_target is None else str(raw_target).strip()
        if not target:
            workbook.close()
            raise ValueError(f"Missing Target for ID {item_id!r} at Excel row {excel_row}.")
        if item_id in id_targets:
            workbook.close()
            raise ValueError(f"Duplicate ID {item_id!r} at Excel row {excel_row}.")
        id_targets[item_id] = target

    workbook.close()
    if not id_targets:
        raise ValueError("No IDs were found in the metadata sheet.")
    return id_targets


def read_csv_sample(path: Path) -> tuple[str, str]:
    last_error: UnicodeDecodeError | None = None
    for encoding in ENCODINGS:
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                return encoding, handle.read(65536)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"Could not decode CSV file {path}: {last_error}")


def detect_dialect(sample: str) -> type[csv.Dialect]:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def inspect_csv(path: Path, relative_path: Path, target: str) -> CsvInfo:
    encoding, sample = read_csv_sample(path)
    dialect = detect_dialect(sample)
    with path.open("r", encoding=encoding, newline="") as handle:
        reader = csv.reader(handle, dialect=dialect)
        try:
            headers = next(reader)
        except StopIteration as exc:
            raise ValueError(f"CSV file is empty: {path}") from exc

    stripped_to_original: dict[str, str] = {}
    for header in headers:
        stripped = header.strip()
        if stripped in stripped_to_original:
            raise ValueError(f"Duplicate CSV column after trimming whitespace in {path}: {stripped!r}")
        stripped_to_original[stripped] = header

    if target in stripped_to_original:
        actual_target = stripped_to_original[target]
    elif "Target" in stripped_to_original:
        # Compatibility fallback for CSVs that use a literal Target column.
        actual_target = stripped_to_original["Target"]
    else:
        raise ValueError(
            f"CSV {path} has no column {target!r} (from Excel Target) "
            "and no fallback column named 'Target'."
        )

    return CsvInfo(path, relative_path, encoding, dialect, actual_target)


def parse_number(text: str, path: Path, row_number: int) -> float | None:
    stripped = text.strip()
    if stripped == "":
        return None
    try:
        value = Decimal(stripped)
    except InvalidOperation as exc:
        raise ValueError(
            f"Non-numeric target value {text!r} in {path}, CSV row {row_number}."
        ) from exc
    if not value.is_finite():
        raise ValueError(f"Non-finite target value {text!r} in {path}, CSV row {row_number}.")
    converted = float(value)
    if not math.isfinite(converted):
        raise ValueError(
            f"Target value {text!r} is outside the supported floating-point range "
            f"in {path}, CSV row {row_number}."
        )
    return converted


def format_number(value: float) -> str:
    if not math.isfinite(value):
        raise ValueError(f"Noise generation produced a non-finite value: {value!r}")
    if value == 0:
        return "0"
    return format(value, ".17g")


def stable_rng(seed: int, noise: Decimal, info: CsvInfo) -> random.Random:
    """Return an order-independent RNG for one noise level and source file."""
    source_key = f"{info.source.parent.name}/{info.relative_to_id.as_posix()}"
    canonical_noise = format(noise.normalize(), "f")
    key = f"{seed}\0{canonical_noise}\0{source_key}".encode("utf-8")
    derived_seed = int.from_bytes(hashlib.sha256(key).digest()[:16], "big")
    return random.Random(derived_seed)


def noise_scale(values: list[float], amplitude: float, mode: str) -> float:
    if mode == "absolute":
        return amplitude
    if mode == "dataset-std":
        if len(values) < 2:
            return 0.0
        mean = math.fsum(values) / len(values)
        variance = math.fsum((value - mean) ** 2 for value in values) / len(values)
        return amplitude * math.sqrt(variance)
    raise ValueError(f"Unsupported fixed noise scale mode: {mode}")


def write_noisy_csv(
    info: CsvInfo,
    destination: Path,
    noise: Decimal,
    mode: str,
    seed: int,
) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    amplitude = float(noise)
    if not math.isfinite(amplitude):
        raise ValueError(f"Noise amplitude is outside the supported range: {noise}")
    rng = stable_rng(seed, noise, info)

    with info.source.open("r", encoding=info.encoding, newline="") as source_handle:
        reader = csv.DictReader(source_handle, dialect=info.dialect)
        if reader.fieldnames is None:
            raise ValueError(f"CSV file has no header: {info.source}")
        fieldnames = reader.fieldnames
        rows = list(reader)

    parsed_targets = [
        parse_number(row[info.target_column], info.source, row_number)
        for row_number, row in enumerate(rows, start=2)
    ]
    non_empty_targets = [value for value in parsed_targets if value is not None]
    fixed_scale = (
        None if mode == "relative" else noise_scale(non_empty_targets, amplitude, mode)
    )

    with destination.open("w", encoding=info.encoding, newline="") as output_handle:
        writer = csv.DictWriter(
            output_handle,
            fieldnames=fieldnames,
            dialect=info.dialect,
            extrasaction="raise",
        )
        writer.writeheader()
        for row, target in zip(rows, parsed_targets):
            if target is not None:
                if mode == "relative":
                    noisy_target = target * (1.0 + rng.gauss(0.0, amplitude))
                else:
                    assert fixed_scale is not None
                    noisy_target = target + rng.gauss(0.0, fixed_scale)
                row[info.target_column] = format_number(noisy_target)
            writer.writerow(row)


def collect_csv_files(
    data_folder: Path,
    id_targets: dict[str, str],
    dataset_indices: Iterable[int],
) -> tuple[list[CsvInfo], list[str]]:
    files: list[CsvInfo] = []
    warnings: list[str] = []
    indices = tuple(dataset_indices)
    if not indices:
        raise ValueError("The dataset index range is empty.")

    for item_id, target in id_targets.items():
        id_folder = data_folder / item_id
        if not id_folder.is_dir():
            warnings.append(f"ID folder not found: {id_folder}")
            continue
        # Process only the selected directly contained numbered CSV files.
        # Files outside the requested range and files in nested folders are ignored.
        csv_paths = [
            id_folder / f"{dataset_index}.csv"
            for dataset_index in indices
            if (id_folder / f"{dataset_index}.csv").is_file()
        ]
        if not csv_paths:
            warnings.append(
                f"No files named {indices[0]}.csv through {indices[-1]}.csv "
                f"found in: {id_folder}"
            )
            continue
        for csv_path in csv_paths:
            files.append(inspect_csv(csv_path, csv_path.relative_to(id_folder), target))
    return files, warnings


def renumbered_output_path(info: CsvInfo, dataset_start: int) -> Path:
    """Map the selected input range to consecutive output names from 0.csv."""
    try:
        source_index = int(info.relative_to_id.stem)
    except ValueError as exc:
        raise ValueError(f"Dataset filename is not numeric: {info.source.name}") from exc
    output_index = source_index - dataset_start
    if output_index < 0:
        raise ValueError(
            f"Dataset index {source_index} is below --dataset-start={dataset_start}"
        )
    return Path(f"{output_index}.csv")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Add reproducible Gaussian noise to each ID's CSV target column."
    )
    parser.add_argument("metadata_file", type=Path, help="Excel file containing ID and Target columns")
    parser.add_argument("data_folder", type=Path, help="Input root containing one folder per ID")
    parser.add_argument("noise_folder", type=Path, help="Output root for noise-level directories")
    parser.add_argument(
        "--noises",
        default=DEFAULT_NOISES,
        help=(
            "Comma-separated Gaussian standard-deviation amplitudes or labels "
            f"(default: {DEFAULT_NOISES})"
        ),
    )
    parser.add_argument(
        "--noise-mode",
        choices=("relative", "dataset-std", "absolute"),
        default="relative",
        help=(
            "Amplitude interpretation: relative gives target*(1+n*Z); "
            "dataset-std gives target+n*std(target)*Z; absolute gives target+n*Z "
            "(default: relative)"
        ),
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=20260825,
        help="Base random seed for reproducible noise (default: 20260825)",
    )
    parser.add_argument(
        "--dataset-start",
        type=int,
        default=DEFAULT_DATASET_START,
        help=f"First dataset index, inclusive (default: {DEFAULT_DATASET_START})",
    )
    parser.add_argument(
        "--dataset-end",
        type=int,
        default=DEFAULT_DATASET_END,
        help=f"Last dataset index, inclusive (default: {DEFAULT_DATASET_END})",
    )
    parser.add_argument(
        "--sheet",
        default="Sampling design",
        help="Metadata worksheet name (default: Sampling design)",
    )
    parser.add_argument(
        "--strict-missing-ids",
        action="store_true",
        help="Treat missing/empty ID folders as an error instead of a warning",
    )
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        noises = parse_noise_list(args.noises)
        if args.dataset_start < 0:
            raise ValueError("--dataset-start must be non-negative.")
        if args.dataset_end < args.dataset_start:
            raise ValueError("--dataset-end must be greater than or equal to --dataset-start.")
        dataset_indices = range(args.dataset_start, args.dataset_end + 1)
        if not args.metadata_file.is_file():
            raise ValueError(f"Metadata file not found: {args.metadata_file}")
        if not args.data_folder.is_dir():
            raise ValueError(f"Data folder not found: {args.data_folder}")

        id_targets = load_id_targets(args.metadata_file, args.sheet)
        csv_files, warnings = collect_csv_files(
            args.data_folder,
            id_targets,
            dataset_indices,
        )

        for warning in warnings:
            print(f"Warning: {warning}", file=sys.stderr)
        if args.strict_missing_ids and warnings:
            raise ValueError("Missing or empty ID folders were found; no files were written.")
        if not csv_files:
            raise ValueError("No CSV files were available to process.")

        # All CSV headers are validated before any output is written.
        written = 0
        for noise in noises:
            level_folder = args.noise_folder / noise_directory_name(noise)
            for info in csv_files:
                item_id = info.source.relative_to(args.data_folder).parts[0]
                destination = (
                    level_folder
                    / item_id
                    / renumbered_output_path(info, args.dataset_start)
                )
                write_noisy_csv(
                    info,
                    destination,
                    noise,
                    args.noise_mode,
                    args.seed,
                )
                written += 1

        print(
            f"Done: {len(csv_files)} source CSV file(s) x {len(noises)} noise level(s) "
            f"= {written} output file(s)."
        )
        print(f"Dataset range: {args.dataset_start}.csv through {args.dataset_end}.csv")
        print(
            "Output numbering: 0.csv through "
            f"{args.dataset_end - args.dataset_start}.csv"
        )
        print(f"Noise mode: {args.noise_mode}; seed: {args.seed}")
        for noise in noises:
            print(f"  {noise} -> {args.noise_folder / noise_directory_name(noise)}")
        return 0
    except (OSError, ValueError, csv.Error) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
